"""Export the three sheets build_data.py reads, so the page can be rebuilt from the repo.

    python tools/export_workbook.py

WHY, and why not just commit the .xlsx. `source/International Property.xlsx` is
gitignored because the repo is public and publishing Charlie's research file is his call.
The consequence was that a fresh clone could serve the page but not regenerate its data:
the chain stopped at a file one person has, on a page whose whole claim is that the
numbers are checkable.

Committing the workbook would fix that and cost two things. It is a binary, so no change
to it is ever reviewable in a diff. And it carries two sheets the build does not read -
"Scoring Matrix" and "AUD Return Scenarios" - the second of which models a 0.5% tax drag,
which is the exact error this whole project exists to correct: for an Australian resident
the drag is the marginal rate, because the foreign tax offset is capped at the foreign
tax paid, so the foreign rate is a floor and the home rate is the bill. Publishing that
sheet alongside the page invites someone to read a superseded model as the page's claim.

So this exports the three sheets that ARE the build's input - Comparison, Country
Profiles, Sources - as CSV. Text, diffable, complete: build_data.py falls back to these
when the workbook is absent, so `git clone && python build_data.py` reproduces data.json
byte for byte. The two superseded sheets stay out, deliberately.
"""

import csv
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "source" / "International Property.xlsx"
DEST = ROOT / "workbook"

# Exactly what build_data.py opens. Adding a sheet here without the build reading it
# would publish something nothing verifies.
SHEETS = {"Comparison": "comparison.csv",
          "Country Profiles": "country_profiles.csv",
          "Sources": "sources.csv"}


def main() -> int:
    if not SRC.exists():
        print(f"no workbook at {SRC}. This script is for the machine that HAS it; "
              f"everyone else reads the CSVs it produced.")
        return 1

    wb = openpyxl.load_workbook(SRC, data_only=True)
    missing = [s for s in SHEETS if s not in wb.sheetnames]
    if missing:
        print(f"workbook is missing {missing}; it has {wb.sheetnames}")
        return 1

    DEST.mkdir(exist_ok=True)
    for sheet, fname in SHEETS.items():
        ws = wb[sheet]
        rows = list(ws.iter_rows(values_only=True))
        # Trailing all-empty rows are an artefact of the spreadsheet's used range and
        # would show up as churn in the diff every time the file is touched.
        while rows and all(c is None or str(c).strip() == "" for c in rows[-1]):
            rows.pop()
        with open(DEST / fname, "w", newline="", encoding="utf-8") as f:
            w = csv.writer(f)
            for r in rows:
                w.writerow(["" if c is None else c for c in r])
        print(f"  {sheet:18} -> workbook/{fname:24} {len(rows):>3} rows")

    skipped = [s for s in wb.sheetnames if s not in SHEETS]
    print(f"\nnot exported, on purpose: {skipped}")
    print("Scoring Matrix and AUD Return Scenarios are superseded models the build "
          "does not read.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
